"""Used once to freeze the tables as printed in the submitted manuscript (main text, .docx) and Supporting
Information (.docx) into reference/paper_tables/ (re-run on the revised R1 files in the revision), so that pipeline/40_verify.py can compare regenerated tables with
the printed ones cell by cell. Header rows are dropped; only data rows are kept, as printed.

Usage: python tools/extract_paper_tables.py <manuscript.docx> <supporting_information.docx>"""
import sys, pathlib, csv
import docx
from docx.oxml.ns import qn

OUT = pathlib.Path(__file__).resolve().parents[1] / "reference" / "paper_tables"; OUT.mkdir(parents=True, exist_ok=True)


def cell_text(tc):
    return " ".join("".join(n.text or "" for n in p.iter() if n.tag in (qn("w:t"), qn("m:t")))
                    for p in tc.iter(qn("w:p"))).strip()


def rows_of(tbl):
    return [[cell_text(tc) for tc in tr.iter(qn("w:tc"))] for tr in tbl._tbl.iter(qn("w:tr"))]


def write(name, rows):
    with open(OUT / f"Table_{name}.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


ms = [t for t in docx.Document(sys.argv[1]).tables if len(t.rows) > 1]     # one row tables are author biographies
assert len(ms) == 5, len(ms)
# main text tables I to IV with their number of header rows; Table V (checklist) carries no result
for name, (i, nh) in {"I": (0, 1), "II": (1, 2), "III": (2, 1), "IV": (3, 2)}.items():
    write(name, rows_of(ms[i])[nh:])
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
from csp.si_revision import SI_NUM
si = docx.Document(sys.argv[2]).tables
assert len(si) == 26, len(si)          # printed in order of first citation (reference/si_order.json)
heads = {SI_NUM[f"S{i + 1}"]: n for i, n in enumerate([1, 2, 2, 2, 2, 1, 1])}          # header rows of the original seven
for k in range(len(si)):
    name = f"S{k + 1}"; rows = rows_of(si[k])[heads.get(name, 1):]
    if name == SI_NUM["S19"]: continue                               # literature table: text, no result
    if name == SI_NUM["S21"]: rows = [[r[0], r[2], r[3]] for r in rows]          # feature table: the definition column is text
    write(name, rows)
print("written", sorted(p.name for p in OUT.glob("*.csv")))
